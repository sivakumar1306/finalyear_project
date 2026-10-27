"""Fill the Review-III template with the project content.

Every new slide reuses the template frame exactly: white background, navy
Calibri 28 pt title at the top left, the VIT logo at the top right and a grey
10 pt slide number at the bottom right.
"""
import copy
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TEMPLATE, LOGO, IMG, REPORT, OUT = sys.argv[1:6]
IMG = Path(IMG)
sys.path.insert(0, REPORT)
from ch_results import REFS  # noqa: E402  (same IEEE list as the report)

NAVY = RGBColor(0x1D, 0x2F, 0x82)
TINT = RGBColor(0xEE, 0xF1, 0xF9)      # light navy card fill
TINT_LINE = RGBColor(0xD3, 0xD9, 0xEE)
ORANGE = RGBColor(0xD9, 0x6C, 0x1E)    # one sharp accent for "ours"
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Calibri'

prs = Presentation(TEMPLATE)
layout = prs.slides[0].slide_layout
N = [0]


# ── basic helpers ───────────────────────────────────────────────────────────
def style_run(r, size, bold=False, color=BLACK, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color


def add_runs(par, text, size, bold=False, color=BLACK, italic=False):
    """**bold** and *italic* inline markup."""
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('*') and len(part) > 1:
            part, i = part[1:-1], True
        r = par.add_run()
        r.text = part
        style_run(r, size, b, color, i)


def set_bullet(par, level=0, char='•'):
    pPr = par._p.get_or_add_pPr()
    mar = 177800 + level * 228600
    pPr.set('marL', str(mar))
    pPr.set('indent', str(-177800))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buSzPct'):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    sz = etree.SubElement(pPr, qn('a:buSzPct')); sz.set('val', '100000')
    bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', char)


def no_bullet(par):
    pPr = par._p.get_or_add_pPr()
    pPr.set('marL', '0'); pPr.set('indent', '0')
    etree.SubElement(pPr, qn('a:buNone'))


def text(slide, x, y, w, h, paras, size=16, anchor='t', align=None, name=None):
    """paras: list of str or (str, opts) — opts: bullet, level, bold, color, size, italic, after, align."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for item in paras:
        s, o = (item, {}) if isinstance(item, str) else item
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if o.get('bullet'):
            set_bullet(par, o.get('level', 0))
        else:
            no_bullet(par)
        par.space_after = Pt(o.get('after', 6))
        a = o.get('align', align)
        if a:
            par.alignment = {'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'l': PP_ALIGN.LEFT}[a]
        add_runs(par, s, o.get('size', size), o.get('bold', False), o.get('color', BLACK), o.get('italic', False))
    return tb


def card(slide, x, y, w, h, fill=TINT, line=None, radius=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def badge(slide, x, y, label, d=0.42, fill=NAVY, size=14):
    """The deck's motif: a navy circle with a white number or letter."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par = tf.paragraphs[0]; par.alignment = PP_ALIGN.CENTER
    r = par.add_run(); r.text = str(label); style_run(r, size, True, WHITE)
    return sh


def arrow(slide, x, y, w=0.35, h=0.3):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = TINT_LINE
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def picture(slide, path, x, y, w=None, h=None, name=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    return pic


def fit_picture(slide, path, x, y, w, h, name=None):
    """Largest picture that fits the box, centred in it."""
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return picture(slide, path, x + (w - pw) / 2, y + (h - ph) / 2, pw, ph, name)


def table(slide, x, y, w, rows, widths, size=13, header=True, row_h=0.36, bold_first_col=False,
          highlight_row=None, align=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for k in ('bandRow', 'firstRow'):
        tblPr.set(k, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # "No Style, Table Grid"
    for ci, cw in enumerate(widths):
        tbl.columns[ci].width = Inches(cw)
    for ri in range(nr):
        tbl.rows[ri].height = Inches(row_h)
        for ci in range(nc):
            c = tbl.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and ri == 0
            c.fill.solid()
            if is_head:
                c.fill.fore_color.rgb = NAVY
            elif highlight_row is not None and ri == highlight_row:
                c.fill.fore_color.rgb = RGBColor(0xFC, 0xEB, 0xDD)
            else:
                c.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF6, 0xF7, 0xFB)
            tf = c.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            a = 'c' if is_head else (align[ci] if align else 'l')
            par.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[a]
            add_runs(par, str(rows[ri][ci]), size, bold=is_head or (bold_first_col and ci == 0),
                     color=WHITE if is_head else BLACK)
            # thin grey grid lines
            tcPr = c._tc.get_or_add_tcPr()
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(side)); ln.set('w', '6350')
                sf = etree.SubElement(ln, qn('a:solidFill'))
                clr = etree.SubElement(sf, qn('a:srgbClr')); clr.set('val', 'C9CEDF')
            # schema order: lines must come before the fill
            fill = tcPr.find(qn('a:solidFill'))
            if fill is not None:
                tcPr.remove(fill); tcPr.append(fill)
    return gt


def bar_chart(slide, x, y, w, h, cats, series, horizontal=False, colors=None, point_colors=None,
              number_format='0.000', max_v=None, title=None, legend=False, label_size=12):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(label_size)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(label_size)
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        style_run(r, 14, True, NAVY)
    else:
        ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size)
    dl.font.color.rgb = BLACK
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE6, 0xEF)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(label_size - 1)
    va.tick_labels.font.color.rgb = GREY
    if max_v is not None:
        va.maximum_scale = max_v
        va.minimum_scale = 0
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(label_size)
    ca.tick_labels.font.color.rgb = BLACK
    ca.format.line.color.rgb = RGBColor(0xBF, 0xC4, 0xD6)
    colors = colors or [NAVY, ORANGE]
    for si, s in enumerate(plot.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = colors[si % len(colors)]
        if point_colors and si == 0:
            for pi, pc in enumerate(point_colors):
                pt = s.points[pi]
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb = pc
    return gf


NOTES_BODY = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr>'
    '<p:cNvPr id="3" name="Notes Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>'
    '<a:p><a:endParaRPr lang="en-US" dirty="0"/></a:p></p:txBody></p:sp>')


def set_notes(slide, txt):
    """The template's notes master has no body placeholder, so add the one its own notes slides use."""
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(etree.fromstring(NOTES_BODY))
    ns.notes_text_frame.text = txt


def new_slide(title, notes=''):
    s = prs.slides.add_slide(layout)
    # template background is plain white
    N[0] += 1
    t = text(s, 0.6, 0.35, 10.5, 0.75, [(title, {'bold': True, 'color': NAVY, 'size': 28, 'after': 0})],
             anchor='m', name='Title')
    picture(s, LOGO, 11.33, 0.35, 1.5, 0.57, name='Logo')
    if notes:
        set_notes(s, notes)
    return s


def caption(s, x, y, w, txt, align='c'):
    text(s, x, y, w, 0.3, [(txt, {'size': 11, 'italic': True, 'color': GREY, 'after': 0})], align=align)


# ═════════════════════════════════════════════════════════════════════════════
# Slide 1 — title (template slide, filled in)
# ═════════════════════════════════════════════════════════════════════════════
s1 = prs.slides[0]
sh = {x.name: x for x in s1.shapes}


def set_box(shape, paras):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    proto = copy.deepcopy(p0._p)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, (s, o) in enumerate(paras):
        par = p0 if i == 0 else None
        if par is None:
            newp = copy.deepcopy(proto)
            for r in newp.findall(qn('a:r')):
                newp.remove(r)
            tf._txBody.append(newp)
            par = tf.paragraphs[-1]
        add_runs(par, s, o['size'], o.get('bold', False), o.get('color', BLACK))


sh['Text 3'].top = Inches(1.40); sh['Text 3'].height = Inches(1.15)
set_box(sh['Text 3'], [('Leveraging Agentic AI for Intelligent Personalized Health Monitoring',
                        {'size': 30, 'bold': True, 'color': NAVY})])
sub = text(s1, 0.6, 2.58, 12.13, 0.4, [('Explainable, Safety-Aware Fall-Risk Prediction from Smart-Ring Data',
                                       {'size': 18, 'color': NAVY, 'after': 0})], align='c', name='Subtitle')
# two students need two lines: grow that row and move the rows below it down
shift = 0.40
sh['Text 5'].height = Inches(0.82)
set_box(sh['Text 5'], [('G Sivakumar — 23BCE1834', {'size': 14}), ('Kriti Sahana V — 23BCE1868', {'size': 14})])
for k in ('Shape 6', 'Text 7', 'Text 8', 'Shape 9', 'Text 10', 'Text 11', 'Shape 12', 'Text 13', 'Text 14',
          'Shape 15'):
    sh[k].top = sh[k].top + Inches(shift)
set_box(sh['Text 8'], [('Dr. Sheena Christabel Pravin, SENSE', {'size': 14})])
set_box(sh['Text 11'], [('B.Tech ECM — SENSE, VIT Chennai', {'size': 14})])
sh['Shape 16'].top = sh['Shape 16'].top + Inches(0.2)
sh['Shape 16'].height = sh['Shape 16'].height + Inches(0.2)
for k in ('Text 17', 'Text 18', 'Text 19'):
    sh[k].top = sh[k].top + Inches(0.3)
# the template's instruction note is not part of the deck
sh['Text 20']._element.getparent().remove(sh['Text 20']._element)
set_box(sh['Text 0'], [('1', {'size': 10, 'color': RGBColor(0x80, 0x80, 0x80)})])
set_notes(s1, 'Introduce the team and the title. This slide is signed by the guide and scanned before the review.')
N[0] = 1

# ═════════════════════════════════════════════════════════════════════════════
# 1. PROJECT OVERVIEW
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('Project Overview — Introduction & Background',
              'Start with the problem: falls are the leading cause of injury in older adults. Wearables detect a '
              'fall only after it happens, LLM chatbots invent numbers, and ML risk scores are black boxes. '
              'Then introduce the three parts we built.')
text(s, 0.7, 1.4, 5.6, 0.4, [('Why this problem matters', {'bold': True, 'color': NAVY, 'size': 20})])
probs = [
    ('Falls are the leading cause of injury-related harm in older adults',
     'Risk rises quietly over days — poor sleep, low HRV, less walking'),
    ('Wearables mostly detect a fall after it has happened',
     'They do not say that today\'s risk is higher than usual'),
    ('LLM health chatbots state numbers they cannot be trusted with',
     'A wrong heart rate or a missed emergency is dangerous'),
]
for i, (a, b) in enumerate(probs):
    y = 2.0 + i * 1.5
    card(s, 0.7, y, 5.6, 1.3)
    badge(s, 0.9, y + 0.2, i + 1)
    text(s, 1.5, y + 0.15, 4.6, 1.05, [(a, {'bold': True, 'size': 15, 'after': 4}), (b, {'size': 14, 'color': GREY})])
text(s, 6.9, 1.4, 5.7, 0.4, [('What we built', {'bold': True, 'color': NAVY, 'size': 20})])
parts = [
    ('Smart-ring app', 'Flutter + Kotlin app collects HR, HRV, SpO₂, temperature, sleep, stress and steps all day'),
    ('Fall-risk engine', 'Turns the vitals into a daily Fall Risk Score (0–100); every point traced to a measurement'),
    ('Verified AI agent', 'Explains the score in plain language, detects emergencies, checks every number it says'),
]
for i, (a, b) in enumerate(parts):
    y = 2.0 + i * 1.5
    card(s, 6.9, y, 5.7, 1.3, fill=WHITE, line=TINT_LINE)
    badge(s, 7.1, y + 0.2, chr(65 + i), fill=ORANGE)
    text(s, 7.7, y + 0.15, 4.7, 1.05, [(a, {'bold': True, 'size': 15, 'after': 4}), (b, {'size': 14, 'color': GREY})])

s = new_slide('Project Overview — Significance & Motivation',
              'Four reasons this matters: predicting instead of reacting, trust through explanations, safety, and '
              'fairness for women. Stress that the system is a screening aid, never a diagnosis.')
mot = [
    ('Predict, don\'t just detect', 'A daily risk estimate gives time to act — rest, review medication, '
     'avoid risky activity — before a fall happens.'),
    ('Trust through transparency', 'The score is a fixed formula with published weights; the app shows how '
     'every point was earned. No black box.'),
    ('Safety first', 'Emergencies are detected from three independent signals, including the ring\'s own fall '
     'event, so a calm-sounding message after a fall is not missed.'),
    ('Fair for women', 'Normal menstrual-cycle changes in temperature, heart rate and HRV are not mistaken for '
     'rising risk.'),
]
for i, (a, b) in enumerate(mot):
    cx, cy = 0.7 + (i % 2) * 6.05, 1.45 + (i // 2) * 2.15
    card(s, cx, cy, 5.85, 1.95)
    badge(s, cx + 0.25, cy + 0.25, i + 1)
    text(s, cx + 0.85, cy + 0.22, 4.8, 1.6, [(a, {'bold': True, 'size': 17, 'color': NAVY, 'after': 6}),
                                            (b, {'size': 15})])
text(s, 0.7, 6.0, 11.9, 0.6, [('Low-cost consumer ring + phone  ·  explainable AI  ·  a screening aid, never a '
                               'diagnosis', {'size': 16, 'italic': True, 'color': NAVY})], align='c')

# ═════════════════════════════════════════════════════════════════════════════
# 2. PROBLEM STATEMENT & OBJECTIVES
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('Problem Statement & Objectives',
              'Read the problem statement, then walk through the objectives table: each objective, how it was '
              'achieved, and the measured outcome.')
card(s, 0.7, 1.35, 11.9, 1.25)
text(s, 0.95, 1.45, 11.4, 1.1, [
    ('Problem statement', {'bold': True, 'color': NAVY, 'size': 16, 'after': 3}),
    ('Given daily vital-sign and activity data from a smart ring, estimate each user\'s fall risk for the next '
     '24 hours in a way that is **explainable**, **personal**, **fair across the menstrual cycle**, and explained by an '
     'AI agent that **never invents numbers** and **never misses an emergency**.', {'size': 15, 'after': 0})])
table(s, 0.7, 2.8, 11.9, [
    ['Objective', 'How it was achieved', 'Outcome'],
    ['Collect continuous vitals from a ring', 'Android app, background measurement loop, Supabase sync', 'Working on our ring'],
    ['Estimate daily fall risk transparently', 'Personal baselines + evidence-weighted engine', 'Every point traceable'],
    ['Explain every score', '"How your score is calculated" page in the app', 'Measure-by-measure breakdown'],
    ['Detect emergencies reliably', 'Keyword + LLM + ring fall-event fusion', '3/3 caught, 0 false alarms'],
    ['Stop the AI inventing numbers', 'Grounded generation + numeric verification', 'Grounding 0.979 (vs 0.723)'],
    ['Treat women fairly', 'Cycle-phase-matched baselines', '43% less false physiology risk'],
    ['Validate on real people', 'PhysioNet LTMM, 73 older adults', 'AUC 0.686 ≈ clinical TUG 0.683'],
], [3.7, 4.6, 3.6], size=13, row_h=0.43, bold_first_col=True)

# ═════════════════════════════════════════════════════════════════════════════
# 3. SYSTEM ARCHITECTURE
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('System Architecture — Overall Block Diagram',
              'Left to right: the ring sends vitals over BLE to the Android app, which stores them on the phone '
              '(Room) and syncs to Supabase. The FastAPI backend on the laptop reads Supabase, runs the fall-risk '
              'engine and the agent, and answers the app. Key principle: the LLM never computes the score.')
fit_picture(s, IMG / 'system_overview.png', 0.7, 1.4, 11.9, 4.0, name='Block diagram')
caption(s, 0.7, 5.42, 11.9, 'Smart ring → Android app → Supabase → agent backend (laptop) → Groq LLM')
card(s, 0.7, 5.85, 11.9, 0.85)
text(s, 0.95, 5.9, 11.4, 0.75, [('**Design principle:** a deterministic engine computes the score; the LLM only '
                                 'explains numbers it is given, and every number is checked. Reproducible and safe.',
                                 {'size': 15, 'after': 0})], anchor='m')

s = new_slide('System Architecture — Agent Workflow',
              'Each chat message runs in three stages. Stage 1 runs the router, the LLM safety classifier and the '
              'ring fall-event lookup in parallel and fuses the three safety signals. Stage 2 fetches only the '
              'needed data. Stage 3 runs the engine, the LLM explains, and every number is verified.')
fit_picture(s, IMG / 'architecture_slide.png', 0.7, 1.3, 11.9, 3.55, name='Agent pipeline')
stages = [
    ('Stage 1 — understand', 'Router picks data streams; keyword, LLM and ring fall-event signals are fused'),
    ('Stage 2 — fetch', 'Only the selected streams are read from Supabase (0.88 on average, not 12)'),
    ('Stage 3 — answer & verify', 'Engine computes; LLM explains (fact → rationale → action); numbers checked'),
]
for i, (a, b) in enumerate(stages):
    x = 0.7 + i * 4.0
    card(s, x, 5.05, 3.85, 1.6)
    badge(s, x + 0.2, 5.2, i + 1)
    text(s, x + 0.75, 5.17, 2.95, 1.4, [(a, {'bold': True, 'size': 15, 'color': NAVY, 'after': 3}),
                                         (b, {'size': 13})])

# ═════════════════════════════════════════════════════════════════════════════
# 4. COMPLETE IMPLEMENTATION
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('Implementation — Module 1: Smart-Ring App',
              'Show the two screenshots. Explain the native BLE plugin, the measurement loop, local storage and '
              'sync, and phone-based fall detection. Mention the pairing fix: measurements now start within '
              'seconds of pairing.')
fit_picture(s, IMG / 'app_home.jpg', 0.7, 1.35, 2.35, 5.25, name='Home screenshot')
fit_picture(s, IMG / 'app_ai_assist.jpg', 3.25, 1.35, 2.35, 5.25, name='AI Assist screenshot')
caption(s, 0.7, 6.62, 2.35, 'Home')
caption(s, 3.25, 6.62, 2.35, 'AI Assist')
feats = [
    ('Ring connection', 'Native Kotlin plugin on the vendor BLE SDK; foreground service keeps the ring connected; '
     'Home card shows status, battery % and Pair / Disconnect'),
    ('Measurement loop', 'Steps → HR → SpO₂ → stress → BP → HRV → temperature; 3-min cooldown, 15 min when '
     'healthy; starts within seconds of the ring binding'),
    ('Storage & sync', 'Every reading saved first in Room (works offline), then batched to Supabase every 5 min'),
    ('Fall detection', 'Phone motion sensors (the ring has no accelerometer stream); a dropped phone is rejected'),
    ('Three tabs', 'Home (vitals + detail charts) · AI Assist (chat) · Fall Risk (score + "how it is calculated")'),
]
for i, (a, b) in enumerate(feats):
    y = 1.35 + i * 1.06
    badge(s, 6.0, y + 0.05, i + 1, d=0.38, size=13)
    text(s, 6.55, y, 6.05, 1.0, [(a, {'bold': True, 'size': 15, 'color': NAVY, 'after': 2}), (b, {'size': 13})])

s = new_slide('Implementation — Module 2: Fall-Risk Engine',
              'Walk through the four steps with the HRV example: median 45 ms, usual variation 7.4, today 30 gives '
              'z = 2.0 and strain 60. The median and MAD are used because one bad day cannot distort them. The '
              'dead zone ignores normal day-to-day wobble.')
steps = [
    ('Personal baseline', 'Last 28 days (≥ 7 points):\nnormal = median\nspread = 1.4826 × MAD'),
    ('Deviation z', 'z = (today − median) ÷ spread,\nsign flipped where low is bad\n(HRV, SpO₂, sleep)'),
    ('Strain 0–100', '100 × clip((z − 0.5) ÷ 2.5)\nz ≤ 0.5 → 0 (dead zone)\nz ≥ 3 → 100'),
    ('Four layers', 'Physiology, Recovery,\nMotion, Fall history —\nweighted average of strains'),
]
for i, (a, b) in enumerate(steps):
    x = 0.7 + i * 3.05
    card(s, x, 1.4, 2.7, 2.15)
    badge(s, x + 0.2, 1.55, i + 1)
    text(s, x + 0.75, 1.55, 1.85, 0.45, [(a, {'bold': True, 'size': 15, 'color': NAVY, 'after': 0})], anchor='m')
    text(s, x + 0.2, 2.1, 2.35, 1.4, [(ln, {'size': 13, 'after': 0}) for ln in b.split('\n')])
    if i < 3:
        arrow(s, x + 2.72, 2.32, 0.3, 0.3)
fit_picture(s, IMG / 'strain_function.png', 0.7, 3.75, 5.6, 2.95, name='Strain function')
table(s, 6.6, 3.8, 6.0, [
    ['Layer', 'Inputs (weight inside the layer)'],
    ['Physiology (P)', 'HRV 0.35, resting HR 0.25, SpO₂ 0.25, temp 0.15'],
    ['Recovery (RC)', 'Sleep 0.40, sleep score 0.30, stress 0.30'],
    ['Motion (MO)', 'Daily steps 0.5, near-falls 0.5'],
    ['Fall history (FH)', '0 / 60 / 100 for 0 / 1 / ≥ 2 falls in a year'],
], [1.75, 4.25], size=13, row_h=0.42, bold_first_col=True)
text(s, 6.6, 6.0, 6.0, 0.7, [('**Example:** HRV median 45 ms, spread 7.4, today 30 → z = 2.0 → strain 60',
                              {'size': 14, 'italic': True, 'color': NAVY, 'after': 0})])

s = new_slide('Implementation — Evidence-Derived Weights & Explainability',
              'The layer weights come from published odds ratios. In logistic regression each risk factor\'s '
              'coefficient is ln(OR), so using ln(OR) as weights is using the published model\'s own '
              'coefficients, scaled to sum to 1. We did not train a model on a small dataset. Point at the app '
              'screenshot: the same weights appear on the score page.')
table(s, 0.7, 1.4, 6.6, [
    ['Layer', 'Odds ratio', 'Source', 'λ'],
    ['Fall history', '2.8', 'Deandrea 2010 (74 studies)', '0.402'],
    ['Motion / gait', '2.1', 'Deandrea 2010', '0.290'],
    ['Physiology', '1.73', 'Mol 2019', '0.214'],
    ['Recovery / sleep', '1.27', 'Baker 2024 (SWAN)', '0.093'],
], [1.75, 1.2, 2.75, 0.9], size=14, row_h=0.42, bold_first_col=True, align=['l', 'c', 'l', 'c'])
card(s, 0.7, 3.65, 6.6, 1.25)
text(s, 0.9, 3.72, 6.2, 1.1, [
    ('λ = ln(OR) ÷ Σ ln(OR)', {'bold': True, 'size': 18, 'color': NAVY, 'after': 4, 'align': 'c'}),
    ('FRS = 0.402·FH + 0.290·MO + 0.214·P + 0.093·RC', {'bold': True, 'size': 16, 'after': 0, 'align': 'c'})])
text(s, 0.7, 5.05, 6.6, 1.7, [
    ('Why ln(OR)? In logistic regression each factor\'s coefficient is β = ln(OR) — so these are the published '
     'models\' own coefficients, scaled to sum to 1.', {'bullet': True, 'size': 14, 'after': 4}),
    ('Tiers: **LOW** < 40 · **MODERATE** 40–69 · **HIGH** ≥ 70', {'bullet': True, 'size': 14, 'after': 4}),
    ('Missing data: weights renormalised; confidence **Q** shows coverage', {'bullet': True, 'size': 14,
                                                                             'after': 0})])
fit_picture(s, IMG / 'app_fall_risk.jpg', 7.6, 1.35, 5.0, 5.0, name='Fall Risk screenshots')
caption(s, 7.6, 6.4, 5.0, 'The app shows the same weights and every layer\'s points')

s = new_slide('Implementation — Module 3: Explainable, Safety-Aware Agent',
              'Router: picks only the data needed; when a question is safety-related and the router is unsure, '
              'it fetches more, not less. Safety: any one strong signal triggers emergency mode; the biometric '
              'signal catches a fall even when the message sounds calm. Grounding: every number is checked and '
              'corrected once if wrong.')
cols = [
    ('Confidence-scored router', [
        'LLM picks which of 12 data streams a question needs',
        '"How did I sleep?" → sleep only; "How am I doing?" → all',
        '**Safety widening:** safety question + confidence < 0.75 → fetch fall risk, HR, SpO₂, HRV too',
        'Keyword router as a fallback']),
    ('Tri-modal safety fusion', [
        '**Keyword** signal: "fell", "chest pain"… (ignores "fell asleep")',
        '**LLM** signal: triggers at confidence ≥ 0.7, overrides a keyword at ≥ 0.85',
        '**Biometric** signal: ring / phone fall in the last 30 min',
        'Any one strong signal → emergency reply']),
    ('Grounded answer + verification', [
        'Engine output given to the LLM as literal values',
        'Reply as JSON: facts → rationale → action',
        'Every number in the reply looked up in the user\'s data',
        'Wrong number → one corrective regeneration']),
]
for i, (h, items) in enumerate(cols):
    x = 0.7 + i * 4.0
    card(s, x, 1.4, 3.85, 5.25)
    badge(s, x + 0.2, 1.58, i + 1)
    text(s, x + 0.75, 1.55, 2.95, 0.55, [(h, {'bold': True, 'size': 16, 'color': NAVY, 'after': 0})], anchor='m')
    text(s, x + 0.25, 2.35, 3.4, 4.2, [(t, {'bullet': True, 'size': 16, 'after': 12}) for t in items])

s = new_slide('Implementation — Integration & Performance Optimisation',
              'Integration path from ring to answer. Optimisations: adaptive measurement cooldown, batched '
              'uploads, selective fetching (0.88 streams per answer instead of 12), and high-rate fall sensors '
              'only when risk is elevated. The pairing fix means the first reading arrives in seconds instead of '
              'minutes.')
chain = ['Smart ring', 'Android app', 'Room (phone)', 'Supabase', 'FastAPI + engine', 'Agent → app']
for i, c in enumerate(chain):
    x = 0.7 + i * 2.02
    card(s, x, 1.45, 1.72, 0.75, fill=NAVY)
    text(s, x + 0.05, 1.45, 1.62, 0.75, [(c, {'bold': True, 'size': 14, 'color': WHITE, 'after': 0})],
         anchor='m', align='c')
    if i < len(chain) - 1:
        arrow(s, x + 1.74, 1.68, 0.26, 0.28)
text(s, 0.7, 2.3, 11.9, 0.4, [('BLE  ·  local-first storage  ·  5-min batched sync  ·  JWT auth  ·  app ↔ laptop '
                               'over Wi-Fi (server address set in the app)', {'size': 13, 'color': GREY,
                                                                                'after': 0})], align='c')
stats = [
    ('0.88', 'data streams fetched per answer (baseline agent: all 12)'),
    ('3 → 15 min', 'adaptive measurement cooldown when vitals are normal'),
    ('5 min', 'batched cloud uploads instead of one request per reading'),
    ('seconds', 'to the first reading after pairing (was a full cooldown)'),
]
for i, (big, small) in enumerate(stats):
    x = 0.7 + i * 3.0
    card(s, x, 3.0, 2.8, 1.9)
    text(s, x + 0.15, 3.15, 2.5, 0.75, [(big, {'bold': True, 'size': 30, 'color': ORANGE, 'after': 0})],
         anchor='m', align='c')
    text(s, x + 0.2, 3.95, 2.4, 0.9, [(small, {'size': 13, 'after': 0})], align='c')
text(s, 0.7, 5.15, 11.9, 1.5, [
    ('High-rate phone motion sensors run only while the risk tier is elevated, then switch off after 10 minutes',
     {'bullet': True, 'size': 15, 'after': 6}),
    ('Router, safety classifier and fall-event lookup run in parallel, so safety adds checks without serial waits',
     {'bullet': True, 'size': 15, 'after': 6}),
    ('CI on GitHub Actions: analyse, test and build the APK on every change',
     {'bullet': True, 'size': 15, 'after': 0})])

# ═════════════════════════════════════════════════════════════════════════════
# 5. RESULTS & TESTING
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('Results — Agent Evaluation (11 Scenarios)',
              'Four versions on the same 11 scenarios, 3 of them emergencies. Only version D caught all three — '
              'one emergency was visible only through the ring\'s fall event, because the message sounded calm. '
              'D also had the highest grounding and was the only one to disclose stale data.')
for i, (big, small) in enumerate([('3 / 3', 'emergencies caught'), ('0 / 8', 'false alarms'),
                                  ('0.979', 'strict numeric grounding')]):
    x = 0.7 + i * 2.05
    card(s, x, 1.4, 1.9, 1.45)
    text(s, x + 0.05, 1.5, 1.8, 0.7, [(big, {'bold': True, 'size': 28, 'color': ORANGE, 'after': 0})],
         anchor='m', align='c')
    text(s, x + 0.1, 2.2, 1.7, 0.6, [(small, {'size': 13, 'after': 0})], align='c')
table(s, 0.7, 3.1, 6.0, [
    ['Version', 'Emergencies', 'Grounding', 'Stale data'],
    ['A – Plain LLM', '0/3', '0.723', '0/1'],
    ['B – Plain RAG', '0/3', '0.606', '0/1'],
    ['C – Baseline agent', '2/3', '0.955', '0/1'],
    ['D – Our agent', '3/3', '0.979', '1/1'],
], [2.1, 1.35, 1.35, 1.2], size=14, row_h=0.42, highlight_row=4, align=['l', 'c', 'c', 'c'])
text(s, 0.7, 5.35, 6.0, 1.3, [
    ('No version raised a false alarm, including the "I fell asleep" keyword trap', {'bullet': True, 'size': 14,
                                                                                      'after': 4}),
    ('Self-verification fired once and raised that reply from 0.80 to 1.00', {'bullet': True, 'size': 14,
                                                                              'after': 0})])
bar_chart(s, 7.0, 1.35, 5.6, 5.3, ['A – Plain LLM', 'B – Plain RAG', 'C – Baseline', 'D – Ours'],
          [('Strict grounding', (0.723, 0.606, 0.955, 0.979))], max_v=1.1,
          point_colors=[NAVY, NAVY, NAVY, ORANGE], title='Strict numeric grounding (higher is better)')

s = new_slide('Results — Fall Risk Scores & Cycle-Aware Baselines',
              'The engine puts each demo user in the expected tier and explains why — for Lakshmi, past falls and '
              'near-falls dominate. On normal luteal days, cycle-aware baselines cut physiology risk by about '
              '40% and removed Meera\'s two false flags, while her genuine strain was still detected.')
table(s, 0.7, 1.4, 6.4, [
    ['Demo user', 'FRS', 'Tier', 'Top contributors'],
    ['Arun (healthy)', '6', 'LOW', 'steps +5.2, SpO₂ +1.1'],
    ['Meera (luteal + poor recovery)', '47', 'MODERATE', 'past fall +24.1, HRV +7.5'],
    ['Ravi (4-day-old data)', '3', 'LOW', 'confidence Q 70, disclosed'],
    ['Lakshmi (recent fall)', '71', 'HIGH', 'falls +40.2, near-falls +14.5'],
    ['Priya (healthy luteal)', '4', 'LOW', 'HRV +2.5'],
], [2.55, 0.65, 1.15, 2.05], size=12, row_h=0.5, align=['l', 'c', 'c', 'l'])
text(s, 0.7, 4.6, 6.4, 2.0, [
    ('Every score is the exact sum of its contributions (checked by a unit test)', {'bullet': True, 'size': 14,
                                                                                     'after': 6}),
    ('Stale data lowers confidence Q, and the agent must say so', {'bullet': True, 'size': 14, 'after': 6}),
    ('Cycle-aware baselines: luteal +0.3 °C, +2 bpm, −7% HRV are treated as normal', {'bullet': True, 'size': 14,
                                                                                      'after': 0})])
bar_chart(s, 7.4, 1.35, 5.2, 4.75, ['Priya (healthy)', 'Meera (strained)'],
          [('Without cycle-aware baseline', (18.1, 19.3)), ('With cycle-aware baseline', (10.2, 11.5))],
          number_format='0.0', max_v=25, legend=True, colors=[RGBColor(0x9A, 0xA6, 0xCF), NAVY],
          title='Physiology risk on normal luteal days')
text(s, 7.4, 6.15, 5.2, 0.5, [('**−43%** (Priya) and **−40%** (Meera); Meera\'s false flags 2 → 0',
                               {'size': 14, 'color': NAVY, 'after': 0})], align='c')

s = new_slide('Results — Validation on Real Data (PhysioNet LTMM)',
              'Our ring users have never fallen, so we validated the approach on PhysioNet LTMM: 73 older adults, '
              '35 fallers and 38 non-fallers. With no training, our gait index scored AUC 0.686, the same as the '
              'Timed Up and Go test doctors use. A trained model reached only 0.721. Limitations: back-worn '
              'sensor, motion layer only, past falls, small sample.')
text(s, 0.7, 1.4, 5.3, 2.95, [
    ('**Why LTMM:** our ring users have no recorded falls; LTMM is the public dataset with real older adults, '
     'movement sensors and known fall history', {'bullet': True, 'size': 14, 'after': 5}),
    ('**Data:** 73 adults (35 fallers, 38 non-fallers), lower-back accelerometer, 1-min lab walk at 100 Hz',
     {'bullet': True, 'size': 14, 'after': 5}),
    ('**Method:** same idea as the engine — no training, fixed rules, feature directions from the literature',
     {'bullet': True, 'size': 14, 'after': 0})])
card(s, 0.7, 4.45, 5.3, 2.2)
text(s, 0.9, 4.52, 4.95, 2.05, [
    ('Honest limitations', {'bold': True, 'size': 15, 'color': NAVY, 'after': 3}),
    ('Sensor on the back, not the finger — validates the approach', {'bullet': True, 'size': 13, 'after': 2}),
    ('Tests the motion layer only', {'bullet': True, 'size': 13, 'after': 2}),
    ('Fallers defined by past falls; n = 73 (CI 0.56–0.81)', {'bullet': True, 'size': 13, 'after': 2}),
    ('3-day home walking did not add information (0.650 vs 0.655)', {'bullet': True, 'size': 13, 'after': 0})])
bar_chart(s, 6.3, 1.35, 6.3, 5.3,
          ['Berg Balance (clinical)', 'Timed Up and Go (clinical)', 'Our index (no training)',
           'Logistic regression (trained)', 'Cadence alone'],
          [('AUC', (0.618, 0.683, 0.686, 0.721, 0.770))], horizontal=True, max_v=1.0,
          point_colors=[NAVY, NAVY, ORANGE, NAVY, NAVY], title='Separating fallers from non-fallers (AUC)')

s = new_slide('Results — Comparison with Existing Solutions',
              'Each existing work solves part of the problem. Our contribution is combining a transparent '
              'evidence-weighted score, cycle-aware baselines, three-signal safety and verified LLM numbers.')
table(s, 0.7, 1.4, 11.9, [
    ['Existing work', 'What it does', 'What is missing', 'Our approach'],
    ['Merrill et al., 2026 (Nature Comms)', 'LLM agent on wearable data, 84% numeric accuracy',
     'No emergency handling', 'Tri-modal safety fusion: 3/3 emergencies'],
    ['Mohammed et al., 2026 (ICAART)', 'Fact → rationale → action health agent', 'Fetches all data; no safety layer',
     'Selective routing (0.88 streams) + safety widening'],
    ['Lu et al., 2026', 'Query-conditioned graph retrieval', 'Retrieval not linked to safety',
     'Safety-asymmetric router'],
    ['SelfCheckGPT, Self-Refine (2023)', 'Hallucination checks, self-critique', 'Not checked against real data',
     'Every number checked against the user\'s data (0.979)'],
    ['Deandrea 2010, Hausdorff 2001', 'Fall risk factors, gait variability', 'Population level, no daily score',
     'Daily personal score using those odds ratios'],
    ['Trained fall models (e.g. TCNs)', 'High accuracy on labelled data', 'Black box, needs fall labels',
     'Untrained, explainable; AUC on par with TUG'],
    ['Wearable baselines (Oura, smartwatch studies)', 'Personal baselines', 'One baseline for the whole cycle',
     'Cycle-phase-matched baselines (−43%)'],
], [3.0, 3.0, 2.75, 3.15], size=13, row_h=0.62, bold_first_col=True)

s = new_slide('Results — Performance Benchmarking & Validation',
              'Latency: in the batch evaluation version D averaged 19.6 s because it makes extra router, safety '
              'and verification calls — the cost of safety. A single live query on the laptop took about 3 s. '
              'Testing: 123 backend and 127 app tests pass, analyser clean, APK builds on every change.')
bar_chart(s, 0.7, 1.35, 5.9, 4.3, ['A – Plain LLM', 'B – Plain RAG', 'C – Baseline', 'D – Ours'],
          [('Mean latency (s)', (2.90, 16.82, 15.98, 19.61))], number_format='0.0', max_v=24,
          point_colors=[NAVY, NAVY, NAVY, ORANGE], title='Mean reply latency in the batch run (s)')
text(s, 0.7, 5.75, 5.9, 0.95, [
    ('Live single query on the laptop: **≈ 3.0 s** (router + safety 1.4 s, fetch, answer + verify 1.6 s)',
     {'size': 14, 'after': 0, 'color': NAVY})])
table(s, 6.9, 1.4, 5.7, [
    ['Check', 'What it verifies', 'Result'],
    ['Engine unit tests', 'Baselines, strain, cycle, tiers', 'Pass'],
    ['Contribution sum', 'Points add up to the score', 'Pass'],
    ['Safety tests', '"fell asleep" guard, override, biometric', 'Pass'],
    ['Grounding tests', 'Digit-boundary matching', 'Pass'],
    ['Backend total', '123 tests', 'All pass'],
    ['App total', '127 tests, incl. 320 px layouts', 'All pass'],
    ['Static analysis', 'flutter analyze', 'No issues'],
    ['Build', 'GitHub Actions APK', 'Success'],
], [1.65, 2.95, 1.1], size=13, row_h=0.47, bold_first_col=True, align=['l', 'l', 'c'])
text(s, 6.9, 5.75, 5.7, 0.95, [('Real-data validation: LTMM AUC 0.686 vs clinical TUG 0.683 (slide 14)',
                                {'size': 14, 'after': 0, 'color': NAVY})])

s = new_slide('Standards Followed & SDG Mapping',
              'Standards that guided design, risk management, data protection and reporting. SDG 3 is the main '
              'goal; SDG 5 through cycle-aware fairness; SDG 9 and 10 through low-cost, explainable technology.')
text(s, 0.7, 1.35, 6.4, 0.4, [('Standards followed', {'bold': True, 'size': 18, 'color': NAVY})])
stds = [
    ('IEC 62304', 'medical software life-cycle — structured development and testing'),
    ('ISO 14971', 'risk management — missed emergencies vs false alarms'),
    ('IEC 82304-1', 'health software safety — screening estimate, never a diagnosis'),
    ('ISO/IEC 23894', 'AI risk management — the LLM never computes the score'),
    ('Bluetooth 5 (BLE), ISO/IEEE 11073', 'ring–phone link, vital-sign data'),
    ('ISO/IEC 27001, DPDP Act 2023', 'protecting biometric data, consent, deletion'),
    ('RFC 8446 (TLS 1.3), RFC 7519 (JWT)', 'encrypted transport, user authentication'),
    ('TRIPOD', 'transparent reporting of the LTMM experiments'),
]
text(s, 0.7, 1.85, 6.4, 4.9, [(f'**{a}** — {b}', {'bullet': True, 'size': 15, 'after': 9}) for a, b in stds])
text(s, 7.4, 1.35, 5.2, 0.4, [('SDG mapping', {'bold': True, 'size': 18, 'color': NAVY})])
sdgs = [
    ('3', RGBColor(0x4C, 0x9F, 0x38), 'Good Health and Well-being', 'Early warning and emergency detection'),
    ('5', RGBColor(0xFF, 0x3A, 0x21), 'Gender Equality', 'Cycle-aware baselines give women fair scores'),
    ('9', RGBColor(0xFD, 0x69, 0x25), 'Industry, Innovation, Infrastructure', 'Ring + cloud + explainable AI'),
    ('10', RGBColor(0xDD, 0x13, 0x67), 'Reduced Inequalities', 'Affordable hardware, plain-language answers'),
]
for i, (n, col, a, b) in enumerate(sdgs):
    y = 1.85 + i * 1.2
    card(s, 7.4, y, 5.2, 1.05)
    badge(s, 7.55, y + 0.2, n, d=0.65, fill=col, size=16)
    text(s, 8.35, y + 0.1, 4.1, 0.9, [(f'SDG {n} — {a}', {'bold': True, 'size': 14, 'after': 2}),
                                      (b, {'size': 13, 'color': GREY, 'after': 0})], anchor='m')

s = new_slide('Live & Video Demonstration',
              'Run the demo on the phone hotspot (campus Wi-Fi often blocks phone-to-laptop traffic). Keep the '
              'laptop terminal visible to show the Version D log. If anything fails, play the recorded video.')
text(s, 0.7, 1.35, 7.0, 0.4, [('Live demonstration plan (≈ 8 min)', {'bold': True, 'size': 18, 'color': NAVY})])
demo = [
    ('Home', 'ring connected with battery %; open a vital card and its chart'),
    ('Pair / Disconnect', 'disconnect, pair again — measurements start within seconds'),
    ('Fall Risk', 'score and tier → "How your score is calculated", one layer\'s points'),
    ('AI Assist', '"How is my overall health?" → real numbers matching Home'),
    ('Safety', '"I fell and feel dizzy" → emergency mode'),
    ('Laptop log', 'router streams, three safety signals, timing, grounding score'),
]
for i, (a, b) in enumerate(demo):
    y = 1.9 + i * 0.78
    badge(s, 0.7, y, i + 1, d=0.42)
    text(s, 1.3, y - 0.02, 6.4, 0.7, [(f'**{a}** — {b}', {'size': 15, 'after': 0})], anchor='m')
card(s, 8.1, 1.4, 4.5, 5.25)
text(s, 8.35, 1.55, 4.0, 5.0, [
    ('Video demonstration', {'bold': True, 'size': 18, 'color': NAVY, 'after': 8}),
    ('Screen recording of the same flow (phone screen + laptop log), kept as a backup and played if the '
     'Wi-Fi or the ring fails during the review.', {'size': 15, 'after': 12}),
    ('Setup', {'bold': True, 'size': 16, 'color': NAVY, 'after': 4}),
    ('Phone and laptop on the phone\'s hotspot', {'bullet': True, 'size': 14, 'after': 3}),
    ('uvicorn main:app --host 0.0.0.0 --port 8000', {'bullet': True, 'size': 14, 'after': 3}),
    ('Laptop IP entered in AI Assist → server', {'bullet': True, 'size': 14, 'after': 0})])

# ═════════════════════════════════════════════════════════════════════════════
# 6. CONCLUSION & FUTURE SCOPE
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide('Conclusion — Project Achievements',
              'A complete, working system from ring to verified answer. Read the four numbers, then the summary.')
for i, (big, small) in enumerate([('3 / 3', 'emergencies caught,\n0 false alarms'),
                                  ('0.979', 'numeric grounding\n(plain LLM 0.723)'),
                                  ('0.686', 'AUC on real data —\nclinical TUG 0.683'),
                                  ('−43%', 'false physiology risk\non normal luteal days')]):
    x = 0.7 + i * 3.0
    card(s, x, 1.4, 2.8, 2.1)
    text(s, x + 0.1, 1.55, 2.6, 0.85, [(big, {'bold': True, 'size': 36, 'color': ORANGE, 'after': 0})],
         anchor='m', align='c')
    text(s, x + 0.15, 2.45, 2.5, 0.95, [(ln, {'size': 14, 'after': 0}) for ln in small.split('\n')], align='c')
text(s, 0.7, 3.8, 11.9, 2.9, [
    ('A complete system: smart ring → Android app → Supabase → explainable engine → verified AI agent, all '
     'modules integrated and working', {'bullet': True, 'size': 16, 'after': 8}),
    ('A transparent Fall Risk Score: personal baselines, evidence-derived weights, every point traceable in the app',
     {'bullet': True, 'size': 16, 'after': 8}),
    ('An agent that never computes the score itself, checks every number, discloses stale data, and catches '
     'emergencies from text and from the ring', {'bullet': True, 'size': 16, 'after': 8}),
    ('An untrained, explainable method that performs on par with a clinical test on real older adults',
     {'bullet': True, 'size': 16, 'after': 0})])

s = new_slide('Limitations & Future Scope',
              'Be upfront about the limitations, then show that each one maps to a concrete next step.')
lim = ['No real fall outcomes from ring users yet (validation used LTMM)',
       'LTMM sensor is on the lower back; the ring gives no raw accelerometer stream',
       'Phone-detected falls are not yet uploaded to the backend',
       'Backend runs on a laptop on the same Wi-Fi',
       'Weights come from published studies, not calibrated on our users']
fut = ['Prospective study: follow users for months and record real falls',
       'Use the ring\'s own motion data for gait features when the SDK allows',
       'Upload phone fall events so the biometric safety signal works on real falls',
       'Deploy the backend to the cloud for use anywhere',
       'Calibrate weights on real outcomes; caregiver alerts with consent']
for i, (h, items, col) in enumerate([('Limitations', lim, NAVY), ('Future scope', fut, ORANGE)]):
    x = 0.7 + i * 6.05
    card(s, x, 1.4, 5.85, 5.25)
    text(s, x + 0.3, 1.55, 5.3, 0.5, [(h, {'bold': True, 'size': 20, 'color': col, 'after': 0})], anchor='m')
    for j, t in enumerate(items):
        y = 2.25 + j * 0.86
        badge(s, x + 0.3, y + 0.05, j + 1, d=0.38, fill=col, size=13)
        text(s, x + 0.85, y, 4.75, 0.8, [(t, {'size': 15, 'after': 0})], anchor='m')

# ═════════════════════════════════════════════════════════════════════════════
# 7. REFERENCES
# ═════════════════════════════════════════════════════════════════════════════
half = 12
for k, chunk in enumerate((REFS[:half], REFS[half:])):
    s = new_slide(f'References ({k + 1}/2)', 'IEEE format; the same references as the report bibliography.')
    start = 1 + k * half
    text(s, 0.7, 1.35, 11.9, 5.4, [(f'[{start + i}]  {r}', {'size': 12, 'after': 6}) for i, r in enumerate(chunk)])

# ── Thank You (template slide 12) ───────────────────────────────────────────
ty = prs.slides[11]

# ═════════════════════════════════════════════════════════════════════════════
# order: title, new slides…, thank you; drop the template's other slides
# ═════════════════════════════════════════════════════════════════════════════
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
keep = [ids[0]] + ids[12:] + [ids[11]]
for e in ids:
    sldIdLst.remove(e)
for e in keep:
    sldIdLst.append(e)

# slide numbers (bottom right, template style) on every slide
for i, sl in enumerate(prs.slides, start=1):
    num = None
    for shp in sl.shapes:
        if shp.has_text_frame and shp.name == 'Text 0':
            num = shp
    if num is not None:
        set_box(num, [(str(i), {'size': 10, 'color': RGBColor(0x80, 0x80, 0x80)})])
    else:
        text(sl, 12.53, 7.05, 0.5, 0.3, [(str(i), {'size': 10, 'color': RGBColor(0x80, 0x80, 0x80), 'after': 0})],
             anchor='m', align='r', name='Slide number')

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
