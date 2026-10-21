"""Fill the B.Tech ECM thesis template with the report content.

Usage: python build.py [pages.json]
  pages.json (optional) maps heading keys -> page labels for the TOC / lists.
"""
import copy
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Pt, Emu

HERE = Path(__file__).parent
TEMPLATE = HERE / 'tpl' / 'template.docx'

sys.path.insert(0, str(HERE))
import content as C  # noqa: E402
OUT = HERE / f"Final_Report_{C.ME[0].title().replace(' ', '_')}.docx"

PAGES = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 else {}
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

doc = docx.Document(TEMPLATE)
body = doc.element.body
els = list(body.iterchildren())

# ── numbering: fresh decimal list instances that restart at 1 ────────────────
numbering = doc.part.numbering_part.element
_next_num = [100]


def new_decimal_list():
    nid = _next_num[0]
    _next_num[0] += 1
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="1"/>'
        f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
    return nid


# ── inline markup: **bold**, *italic*, _{sub}, ^{sup} ──────────────────────────
TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|_\{[^}]*\}|\^\{[^}]*\})')


def add_runs(p, text, size=None, italic=False, bold=False, color='000000', font=None, highlight=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, va = bold, italic, None
        if part.startswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 1:
            part, i = part[1:-1], True
        elif part.startswith('_{'):
            part, va = part[2:-1], 'subscript'
        elif part.startswith('^{'):
            part, va = part[2:-1], 'superscript'
        r = p.add_run(part)
        r.bold = b or None
        r.italic = i or None
        rpr = r._r.get_or_add_rPr()
        if va:
            rpr.append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="{va}"/>'))
        if color:
            rpr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>'))
        if size:
            r.font.size = Pt(size)
        if font:
            r.font.name = font
            rpr.get_or_add_rFonts().set(qn('w:eastAsia'), font)
            rpr.get_or_add_rFonts().set(qn('w:cs'), font)
        if highlight:
            rpr.append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="{highlight}"/>'))
    return p


def ppr(p, xml):
    """Replace the paragraph's pPr with the given inner XML."""
    old = p._p.find(qn('w:pPr'))
    if old is not None:
        p._p.remove(old)
    p._p.insert(0, parse_xml(f'<w:pPr {nsdecls("w")}>{xml}</w:pPr>'))


class Writer:
    """Appends paragraphs before a fixed anchor element."""

    def __init__(self, anchor):
        self.anchor = anchor

    def _p(self):
        p = doc.add_paragraph()
        self.anchor.addprevious(p._p)
        return p

    def body(self, text, after=160, indent_first=False):
        p = self._p()
        ind = '<w:ind w:firstLine="0"/>' if not indent_first else ''
        ppr(p, f'<w:spacing w:after="{after}" w:line="360" w:lineRule="auto"/>{ind}<w:jc w:val="both"/>')
        return add_runs(p, text)

    def front_title(self, text, page_break=True):
        p = self._p()
        pb = '<w:pageBreakBefore/>' if page_break else ''
        ppr(p, f'{pb}<w:spacing w:before="80" w:after="400"/><w:jc w:val="center"/>')
        return add_runs(p, text, size=16, bold=True)

    def chapter(self, num, title):
        p = self._p()
        ppr(p, '<w:keepNext/><w:pageBreakBefore/><w:spacing w:after="80"/><w:jc w:val="center"/>')
        add_runs(p, f'CHAPTER {num}', size=16, bold=True)
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="40" w:after="280"/><w:jc w:val="center"/>')
        add_runs(p, title, size=16, bold=True)

    def h2(self, text):
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="320" w:after="120"/>')
        return add_runs(p, text, size=14, bold=True)

    def h3(self, text):
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="200" w:after="80"/>')
        return add_runs(p, text, bold=True)

    def bullets(self, items):
        for it in items:
            p = self._p()
            ppr(p, '<w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>'
                   '<w:spacing w:before="40" w:after="40" w:line="360" w:lineRule="auto"/><w:jc w:val="both"/>')
            add_runs(p, it)

    def numbered(self, items):
        nid = new_decimal_list()
        for it in items:
            p = self._p()
            ppr(p, f'<w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{nid}"/></w:numPr>'
                   '<w:spacing w:before="40" w:after="40" w:line="360" w:lineRule="auto"/><w:jc w:val="both"/>')
            add_runs(p, it)

    def eq(self, text, num):
        p = self._p()
        ppr(p, '<w:tabs><w:tab w:val="center" w:pos="4333"/><w:tab w:val="right" w:pos="8666"/></w:tabs>'
               '<w:spacing w:before="120" w:after="120"/>')
        r = p.add_run()
        r._r.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
        add_runs(p, text, italic=False)
        r = p.add_run()
        r._r.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
        add_runs(p, f'({num})')

    def caption(self, text, before=60, after=240):
        p = self._p()
        ppr(p, f'<w:keepLines/><w:spacing w:before="{before}" w:after="{after}"/><w:jc w:val="center"/>')
        return add_runs(p, text, size=11, italic=True)

    def table_caption(self, text):
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="200" w:after="80"/><w:jc w:val="center"/>')
        return add_runs(p, text, size=11, italic=True)

    def figure(self, path, caption, width_cm=14.5):
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="160" w:after="60"/><w:jc w:val="center"/>')
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption(caption)

    def picture(self, path, width_cm):
        p = self._p()
        ppr(p, '<w:keepNext/><w:spacing w:before="160" w:after="200"/><w:jc w:val="center"/>')
        p.add_run().add_picture(str(path), width=Cm(width_cm))

    def placeholder(self, text, caption):
        p = self._p()
        ppr(p, '<w:keepNext/><w:pBdr><w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
               '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
               '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
               '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:pBdr>'
               '<w:spacing w:before="160" w:after="60"/><w:jc w:val="center"/>')
        add_runs(p, text, size=11, italic=True, color='888888')
        self.caption(caption)

    def code(self, lines):
        for ln in lines:
            p = self._p()
            ppr(p, '<w:keepLines/><w:spacing w:before="0" w:after="0"/><w:ind w:left="360"/>')
            r = p.add_run(ln if ln else ' ')
            r.font.name = 'Courier New'
            r.font.size = Pt(8)
            rpr = r._r.get_or_add_rPr()
            rpr.get_or_add_rFonts().set(qn('w:eastAsia'), 'Courier New')
            rpr.get_or_add_rFonts().set(qn('w:cs'), 'Courier New')

    def spacer(self, after=0):
        p = self._p()
        ppr(p, f'<w:spacing w:before="0" w:after="{after}"/>')

    def table(self, header, rows, widths_cm, size=10, header_shade='E7E6E6', align=None):
        tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
        self.anchor.addprevious(tbl._tbl)
        tbl.style = doc.styles['Table Grid']
        tbl.alignment = 1
        total = sum(widths_cm)
        tw = int(total * 567)
        tblPr = tbl._tbl.tblPr
        for tag in ('w:tblW', 'w:jc'):
            e = tblPr.find(qn(tag))
            if e is not None:
                tblPr.remove(e)
        style_el = tblPr.find(qn('w:tblStyle'))
        anchor = style_el if style_el is not None else None
        new = [parse_xml(f'<w:tblW {nsdecls("w")} w:w="{tw}" w:type="dxa"/>'),
               parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>')]
        for e in reversed(new):
            if anchor is not None:
                anchor.addnext(e)
            else:
                tblPr.insert(0, e)
        look = tblPr.find(qn('w:tblLook'))
        lay = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
        if look is not None:
            look.addprevious(lay)
        else:
            tblPr.append(lay)
        grid = tbl._tbl.tblGrid
        for gc, wcm in zip(grid.findall(qn('w:gridCol')), widths_cm):
            gc.set(qn('w:w'), str(int(wcm * 567)))
        for ri, row in enumerate([header] + list(rows)):
            tr = tbl.rows[ri]._tr
            trPr = tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if ri == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for ci, val in enumerate(row):
                cell = tbl.rows[ri].cells[ci]
                tcPr = cell._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(old)
                tcPr.insert(0, parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(widths_cm[ci] * 567)}" w:type="dxa"/>'))
                if ri == 0 and header_shade:
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{header_shade}"/>'))
                tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
                p = cell.paragraphs[0]
                jc = 'center' if ri == 0 else ((align[ci] if align else 'left'))
                ppr(p, f'<w:spacing w:before="20" w:after="20"/><w:jc w:val="{jc}"/>')
                add_runs(p, str(val), size=size, bold=(ri == 0))
        # space after table
        self.spacer(after=120)
        return tbl

    def toc_line(self, text, page, indent=0, italic_page=False, before=20):
        p = self._p()
        ind = f'<w:ind w:left="{indent}"/>' if indent else ''
        ppr(p, f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="8666"/></w:tabs>'
               f'<w:spacing w:before="{before}" w:after="20"/>{ind}')
        add_runs(p, text)
        r = p.add_run()
        r._r.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
        add_runs(p, str(page), italic=italic_page)

    def bib(self, text):
        p = self._p()
        ppr(p, '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr>'
               '<w:spacing w:before="60" w:after="120"/><w:jc w:val="both"/>')
        add_runs(p, text)


# ── 1. Front pages kept from the template: title, declaration, certificate ────
def set_text(par_el, text_parts):
    """Replace all runs of a paragraph with [(text, dict(bold=..))] keeping pPr."""
    p = docx.text.paragraph.Paragraph(par_el, doc)
    for r in list(par_el.findall(qn('w:r'))) + list(par_el.findall(qn('w:proofErr'))):
        par_el.remove(r)
    for t, kw in text_parts:
        add_runs(p, t, **kw)


T = C.TITLE
set_text(els[4], [(C.MAIN_TITLE, dict(bold=True, size=18))])
subtitle = copy.deepcopy(els[4]); els[4].addnext(subtitle)
set_text(subtitle, [(C.SUBTITLE, dict(size=14))])
set_text(els[12], [(C.STUDENTS[0][0], {})])
set_text(els[13], [(C.STUDENTS[0][1], {})])
# second student under the first
second = copy.deepcopy(els[12]); els[13].addnext(second)
set_text(second, [(C.STUDENTS[1][0], {})])
second_reg = copy.deepcopy(els[13]); second.addnext(second_reg)
set_text(second_reg, [(C.STUDENTS[1][1], {})])
# drop two spacer lines to keep the title page on one page
for el in (els[14], els[15]):
    body.remove(el)
set_text(els[26], [(C.MONTH_YEAR, {})])

# Declaration
NAME, REG = C.ME
set_text(els[33], [
    ('I hereby declare that the report titled ', {}), (f'“{T}”', dict(bold=True)),
    (f' submitted by me, {NAME} ({REG}), to the School of Electronics Engineering, Vellore Institute of Technology, '
     'Chennai, in partial fulfilment of the requirements for the award of ', {}), ('Bachelor of Technology', dict(bold=True)),
    (' in ', {}), ('Electronics and Computer Engineering', dict(bold=True)),
    (', is a bonafide record of the work carried out by me under the supervision of ', {}),
    (C.GUIDE, dict(bold=True)), ('.', {})])
set_text(els[34], [('I further declare that the work reported in this report has not been submitted and will not be '
                    'submitted, either in part or in full, for the award of any other degree or diploma of this '
                    'institute or of any other institute or University.', {})])
set_text(els[35], [('Signature of the Candidate:', {})])
set_text(els[36], [(f'Name and Reg. No.: {NAME} ({REG})', {})])
# Certificate
set_text(els[56], [
    ('This is to certify that the project report titled ', {}), (f'“{T}”', dict(bold=True)),
    (f' submitted by {NAME} ({REG}) to Vellore '
     'Institute of Technology, Chennai, in partial fulfilment of the requirement for the award of the degree of ', {}),
    ('Bachelor of Technology', dict(bold=True)), (' in ', {}), ('Electronics and Computer Engineering', dict(bold=True)),
    (' is a bonafide work carried out under my supervision. The project report fulfils the requirements as per the '
     'regulations of this University and, in my opinion, meets the necessary standards for submission. The contents '
     'of this report have not been submitted and will not be submitted either in full or in part, for the award of '
     'any other degree or diploma and the same is certified.', {})])

# remove hand-typed page numbers (i, ii) on declaration / certificate and trailing spacers
for idx in (49, 61):
    body.remove(els[idx])
for idx in range(38, 49):  # empty spacers after declaration
    body.remove(els[idx])

# ── 2. Delete template sample content after the certificate ───────────────────
front_sect_p = els[348]            # paragraph carrying the front-matter sectPr
final_sect = els[-1]
for el in els[62:463]:
    if el is front_sect_p:
        continue
    if el.getparent() is not None:
        body.remove(el)

# Sections: (1) title page, (2) front matter in roman numerals, (3) chapters.
fs = front_sect_p.find('.//' + qn('w:sectPr'))
# section 1 break after the title page date paragraph
title_sect = copy.deepcopy(fs)
ppr_el = els[26].find(qn('w:pPr'))
if ppr_el is None:
    ppr_el = parse_xml(f'<w:pPr {nsdecls("w")}/>')
    els[26].insert(0, ppr_el)
ppr_el.append(title_sect)
# remove the empty page-break paragraph that followed the title page (section break does the job)
if els[27].getparent() is not None:
    body.remove(els[27])

# roman footer for the front section
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.opc.part import Part  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
footer1_rel = [r for r in doc.part.rels.values() if r.reltype == RT.FOOTER][0]
ftr_xml = footer1_rel.target_part.blob
footer2 = Part(PackURI('/word/footer2.xml'), footer1_rel.target_part.content_type, ftr_xml, doc.part.package)
rid = doc.part.relate_to(footer2, RT.FOOTER)
fs.insert(0, parse_xml(f'<w:footerReference {nsdecls("w", "r")} w:type="default" r:id="{rid}"/>'))
# explicit empty footer on the title page so no viewer carries a page number onto it
empty_ftr = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             b'<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
             b'<w:p><w:pPr><w:pStyle w:val="Footer"/></w:pPr></w:p></w:ftr>')
footer3 = Part(PackURI('/word/footer3.xml'), footer1_rel.target_part.content_type, empty_ftr, doc.part.package)
rid3 = doc.part.relate_to(footer3, RT.FOOTER)
title_sect.insert(0, parse_xml(f'<w:footerReference {nsdecls("w", "r")} w:type="default" r:id="{rid3}"/>'))
pg = fs.find(qn('w:pgNumType'))
if pg is not None:
    fs.remove(pg)
fs.find(qn('w:pgMar')).addnext(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>'))
# chapters: arabic from 1 (already in final sectPr)

# Certificate starts a new page (template used an empty page-break paragraph)
pp = els[51].find(qn('w:pPr'))
if pp is None:
    pp = parse_xml(f'<w:pPr {nsdecls("w")}/>'); els[51].insert(0, pp)
pp.insert(0, parse_xml(f'<w:pageBreakBefore {nsdecls("w")}/>'))
if els[50].getparent() is not None:
    body.remove(els[50])

# ── 3. Front matter (abstract … acronyms) before the front section break ──────
ch = Writer(final_sect)
C.chapters(ch, PAGES)          # first, so figure/table registries are filled
fm = Writer(front_sect_p)
C.front_matter(fm, PAGES)

doc.save(OUT)
print('saved', OUT)
